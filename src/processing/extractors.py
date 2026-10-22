from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from processing.cleaning import clean_markdown
from processing.models import ExtractedDocument, ExtractedSection

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".md", ".markdown",
    ".csv", ".json", ".html", ".htm", ".txt",
}


def extract_document(path: str | Path, *, enable_ocr: bool = True) -> ExtractedDocument:
    """Extrai um arquivo suportado para seções de texto com localização."""
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Formato não suportado: {ext or '(sem extensão)'}")

    dispatch = {
        ".pdf": lambda: _extract_pdf(file_path, enable_ocr=enable_ocr),
        ".docx": lambda: _extract_docx(file_path),
        ".xlsx": lambda: _extract_excel(file_path),
        ".xlsm": lambda: _extract_excel(file_path),
        ".pptx": lambda: _extract_pptx(file_path),
        ".md": lambda: _extract_markdown(file_path),
        ".markdown": lambda: _extract_markdown(file_path),
        ".csv": lambda: _extract_csv(file_path),
        ".json": lambda: _extract_json(file_path),
        ".html": lambda: _extract_html(file_path),
        ".htm": lambda: _extract_html(file_path),
        ".txt": lambda: _extract_text(file_path),
    }
    sections = dispatch[ext]()
    stat = file_path.stat()
    document_metadata = {
        "source_path": str(file_path.resolve()),
        "file_size_bytes": stat.st_size,
        "file_created_at": _iso_timestamp(stat.st_ctime),
        "file_modified_at": _iso_timestamp(stat.st_mtime),
    }
    document_metadata.update(_extract_internal_metadata(file_path, ext))
    return ExtractedDocument(
        source_file=file_path.name,
        file_type=ext.lstrip("."),
        sections=sections,
        document_metadata=document_metadata,
    )

def _extract_internal_metadata(path: Path, extension: str) -> dict[str, Any]:
    """Lê autor e datas internas quando o formato oferece essas propriedades."""
    properties: Any = None
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            properties = PdfReader(str(path)).metadata
            if properties is None:
                return {}
            return {
                key: value
                for key, value in {
                    "document_author": getattr(properties, "author", None),
                    "document_created_at": _metadata_datetime(getattr(properties, "creation_date", None)),
                    "document_modified_at": _metadata_datetime(getattr(properties, "modification_date", None)),
                }.items()
                if value is not None
            }
        if extension == ".docx":
            from docx import Document
            properties = Document(str(path)).core_properties
        elif extension == ".pptx":
            from pptx import Presentation
            properties = Presentation(str(path)).core_properties
        elif extension in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            try:
                properties = workbook.properties
                result = {
                    "document_author": getattr(properties, "creator", None),
                    "document_created_at": _metadata_datetime(getattr(properties, "created", None)),
                    "document_modified_at": _metadata_datetime(getattr(properties, "modified", None)),
                }
            finally:
                workbook.close()
            return {key: value for key, value in result.items() if value is not None}
        else:
            return {}
    except (ImportError, OSError, ValueError, KeyError):
        # Metadados internos são complementares; a extração principal já ocorreu.
        return {}

    return {
        key: value
        for key, value in {
            "document_author": getattr(properties, "author", None),
            "document_created_at": _metadata_datetime(getattr(properties, "created", None)),
            "document_modified_at": _metadata_datetime(getattr(properties, "modified", None)),
        }.items()
        if value is not None
    }

def _metadata_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    return str(value)


def _iso_timestamp(timestamp: float) -> str:
    from datetime import datetime, timezone

    return datetime.fromtimestamp(timestamp, tz=timezone.utc).isoformat()


def _extract_pdf(path: Path, *, enable_ocr: bool) -> list[ExtractedSection]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections: list[ExtractedSection] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if enable_ocr and len(text) < 20:
            # Não falha silenciosamente: se o OCR não estiver configurado, o
            # consumidor recebe uma mensagem clara em vez de perder conteúdo.
            text = _ocr_pdf_page(path, page_number)
        if text:
            sections.append(ExtractedSection(text, "page", page_number))
    return sections


def _ocr_pdf_page(path: Path, page_number: int) -> str:
    try:
        from pdf2image import convert_from_path
        import pytesseract
    except ImportError as exc:
        raise ImportError("OCR requer pdf2image, Pillow e pytesseract.") from exc

    try:
        images = convert_from_path(
            str(path), first_page=page_number, last_page=page_number, dpi=250
        )
    except Exception as exc:
        raise RuntimeError(
            "Não foi possível rasterizar a página para OCR. Verifique se o Poppler está instalado."
        ) from exc
    if not images:
        return ""
    try:
        return pytesseract.image_to_string(images[0], lang="por+eng").strip()
    except Exception as exc:
        raise RuntimeError(
            "Falha no OCR. Verifique se o Tesseract e os idiomas por/eng estão instalados."
        ) from exc


def _extract_docx(path: Path) -> list[ExtractedSection]:
    from docx import Document

    document = Document(str(path))
    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "") or ""
        if style_name.lower().startswith("heading") or style_name.lower().startswith("título"):
            current_heading = text
            sections.append(ExtractedSection(text, "section", current_heading, current_heading))
        else:
            sections.append(ExtractedSection(text, "section", current_heading or "document", current_heading))
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append(ExtractedSection("\n".join(rows), "table", table_index, f"Tabela {table_index}"))
    return sections


def _extract_excel(path: Path) -> list[ExtractedSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sections: list[ExtractedSection] = []
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(value).strip() if value is not None else f"Coluna {i + 1}" for i, value in enumerate(rows[0])]
            lines: list[str] = []
            for row_number, row in enumerate(rows[1:], start=2):
                pairs = []
                for index, value in enumerate(row):
                    if value is None or str(value).strip() == "":
                        continue
                    header = headers[index] if index < len(headers) and headers[index] else f"Coluna {index + 1}"
                    pairs.append(f"{header}: {value}")
                if pairs:
                    lines.append(f"Linha {row_number} — " + "; ".join(pairs))
            if not lines:
                # Planilhas de uma única linha ou sem cabeçalho também são preservadas.
                if any(value is not None for value in rows[0]):
                    lines.append("Cabeçalhos: " + " | ".join(headers))
            if lines:
                sections.append(ExtractedSection("\n".join(lines), "sheet", worksheet.title, worksheet.title))
    finally:
        workbook.close()
    return sections


def _extract_pptx(path: Path) -> list[ExtractedSection]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[ExtractedSection] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                    if any(values):
                        parts.append(" | ".join(values))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notas do apresentador: " + notes)
        except (AttributeError, KeyError):
            pass
        if parts:
            sections.append(ExtractedSection("\n".join(parts), "slide", slide_number))
    return sections


def _extract_markdown(path: Path) -> list[ExtractedSection]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sections: list[ExtractedSection] = []
    heading: str | None = None
    buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            if buffer:
                body = clean_markdown("\n".join(buffer)).strip()
                if body:
                    sections.append(ExtractedSection(body, "section", heading or "document", heading))
                buffer = []
            heading = clean_markdown(stripped).strip() or heading
            if heading:
                sections.append(ExtractedSection(heading, "section", heading, heading))
        elif stripped:
            buffer.append(line)
        elif buffer:
            body = clean_markdown("\n".join(buffer)).strip()
            if body:
                sections.append(ExtractedSection(body, "section", heading or "document", heading))
            buffer = []
    if buffer:
        body = clean_markdown("\n".join(buffer)).strip()
        if body:
            sections.append(ExtractedSection(body, "section", heading or "document", heading))
    return sections


def _extract_csv(path: Path) -> list[ExtractedSection]:
    sections: list[ExtractedSection] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        sample = file.read(4096)
        file.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(file, dialect))
    if not rows:
        return sections
    headers = rows[0]
    lines: list[str] = []
    for row_number, row in enumerate(rows[1:], start=2):
        pairs = []
        for index, value in enumerate(row):
            if not value.strip():
                continue
            header = headers[index].strip() if index < len(headers) and headers[index].strip() else f"Coluna {index + 1}"
            pairs.append(f"{header}: {value.strip()}")
        if pairs:
            lines.append(f"Linha {row_number} — " + "; ".join(pairs))
    if lines:
        sections.append(ExtractedSection("\n".join(lines), "rows", f"2-{len(rows)}"))
    elif any(value.strip() for value in headers):
        sections.append(ExtractedSection("Cabeçalhos: " + " | ".join(headers), "row", 1))
    return sections


def _extract_json(path: Path) -> list[ExtractedSection]:
    data: Any = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
    text = _json_to_text(data)
    return [ExtractedSection(text, "json_path", "$" )] if text else []


def _json_to_text(value: Any, prefix: str = "") -> str:
    lines: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            key_path = f"{prefix}.{key}" if prefix else str(key)
            if isinstance(item, (dict, list)):
                nested = _json_to_text(item, key_path)
                if nested:
                    lines.append(nested)
            else:
                lines.append(f"{key_path}: {item}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            lines.append(_json_to_text(item, f"{prefix}[{index}]"))
    else:
        lines.append(f"{prefix}: {value}" if prefix else str(value))
    return "\n".join(line for line in lines if line)


def _extract_html(path: Path) -> list[ExtractedSection]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8-sig", errors="replace"), "html.parser")
    for element in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
        element.decompose()
    sections: list[ExtractedSection] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr"]):
        text = element.get_text(" ", strip=True)
        if text:
            kind = "section" if element.name.startswith("h") else "paragraph"
            sections.append(ExtractedSection(text, kind, element.name))
    if not sections:
        text = soup.get_text("\n", strip=True)
        if text:
            sections.append(ExtractedSection(text, "document", None))
    return sections


def _extract_text(path: Path) -> list[ExtractedSection]:
    text = path.read_text(encoding="utf-8-sig", errors="replace").strip()
    return [ExtractedSection(text, "document", None)] if text else []